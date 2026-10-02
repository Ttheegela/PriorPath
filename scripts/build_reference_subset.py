"""Build data/reference/subset from raw CMS downloads (run locally, not in CI).

Example:
  python scripts/build_reference_subset.py \
    --ptp data/reference/raw/ptp_practitioner_f1.xlsx --ptp data/reference/raw/ptp_practitioner_f2.xlsx \
    --mue data/reference/raw/mue_practitioner.xlsx --pfs data/reference/raw/PPRRVU26.csv \
    --codes data/reference/codes.txt \
    --ncci NCCI-2026Q4:2026-10-01:2026-12-31 --mue-version MUE-2026Q4:2026-10-01:2026-12-31 \
    --pfs-version PFS-2026:2026-01-01:2026-12-31 --out data/reference/subset
"""

import argparse
import csv
import io
import sys
from datetime import date
from pathlib import Path

from app.reference.base import INVALID_STATUSES, InMemoryReference, RefVersion
from app.reference.cms_adapters import parse_mue, parse_pfs, parse_ptp
from app.reference.normalized import write_normalized


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_table(path: Path) -> list[list[str]]:
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        return [[_cell(v) for v in r] for ws in wb.worksheets for r in ws.iter_rows(values_only=True)]
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:5000], delimiters=",\t|")
    return list(csv.reader(io.StringIO(text), dialect))


def _version(spec: str, kind: str) -> RefVersion:
    name, start, end = spec.split(":")
    return RefVersion(name, kind, date.fromisoformat(start), date.fromisoformat(end))


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--ptp", type=Path, action="append", required=True)
    p.add_argument("--mue", type=Path, required=True)
    p.add_argument("--pfs", type=Path, required=True)
    p.add_argument("--codes", type=Path, required=True)
    p.add_argument("--ncci", required=True, help="NAME:valid_from:valid_to")
    p.add_argument("--mue-version", required=True)
    p.add_argument("--pfs-version", required=True)
    p.add_argument("--out", type=Path, required=True)
    a = p.parse_args()

    codes = {c.strip().upper() for c in a.codes.read_text().split() if c.strip()}
    ncci, mue_v, pfs_v = (
        _version(a.ncci, "ncci"),
        _version(a.mue_version, "mue"),
        _version(a.pfs_version, "pfs"),
    )

    ptp = [
        e
        for f in a.ptp
        for e in parse_ptp(read_table(f), ncci.ref_version)
        if e.col1 in codes and e.col2 in codes
    ]
    mue = [m for m in parse_mue(read_table(a.mue), mue_v.ref_version) if m.code in codes]
    fees, statuses = parse_pfs(read_table(a.pfs), pfs_v.ref_version)
    fees = [f for f in fees if f.code in codes]
    statuses = [s for s in statuses if s.code in codes]

    write_normalized(InMemoryReference([ncci, mue_v, pfs_v], ptp, mue, fees, statuses), a.out)
    invalid = sum(s.status in INVALID_STATUSES for s in statuses)
    by_ind = {i: sum(e.modifier_indicator == i for e in ptp) for i in "019"}
    print(f"ptp={len(ptp)} {by_ind} mue={len(mue)} fees={len(fees)} codes={len(statuses)} invalid={invalid}")
    if by_ind["0"] < 10 or by_ind["1"] < 10 or invalid < 1:
        print(
            "Subset too thin: need >=10 PTP pairs with indicator 0 and 1, "
            "and >=1 code with status D or I. Add codes.",
            file=sys.stderr,
        )
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
